import re
from typing import Dict, List, Tuple
from openpyxl import Workbook, load_workbook
from openpyxl.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from table_structure import MineChemistryInfo, MineContentInfo, MinePhysicsInfo, MineThermalInfo, MineXRDInfo, SampleInfo


'''
样品号
样品名称 类型
黏土基质
石英粉砂
砂
    石英
    长石
    其他矿物 其他
岩屑 砂或岩屑
空洞
其他矿物
'''

def is_number(num):
    pattern = re.compile(r'^[-+]?[-0-9]\d*\.\d*|[-+]?\.?[0-9]\d*$')
    result = pattern.match(num)
    if result:
        return True
    else:
        return False

def get_range(ws) -> List[Tuple[int, int, int, int]]:
    range_list = []
    step = 0
    c1, c2 = 0, 0
    for i, row in enumerate(ws.rows):
        if row[0].font.b:
            if step:
                range_list.append((step, i, c1, c2))
                c1, c2 = 0, 0
            for j, cell in enumerate(row):
                if isinstance(cell, MergedCell) or (cell.value != None and str(cell.value).strip()):
                    if c1 == 0:
                        c1 = j + 1
                    c2 = j + 1
                elif c2:
                    break
            step = i + 1
    for j, cell in enumerate(row):
        if cell.value != None and str(cell.value).strip():
            if c1 == 0:
                c1 = j + 1
            c2 = j + 1
        elif c2:
            break
    range_list.append((step, i + 1, c1, c2))
    return range_list

def mine_content_process(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int, cover: bool) -> Tuple[Dict[str, MineContentInfo], int]:
    head_map = {
        "clay": 0,
        "debris": 0,
        "hollow": 0,
        "id": 0,
        "other": 0,
        "quartz": 0,
        "sampleName": 0,
        "sand": {
            "Ominerals": 0,
            "feldspar": 0,
            "quartz": 0
        }
    }
    cover_num = 0
    instance_list = {}
    sand_list = []
    for i, row in enumerate(ws.iter_rows(min_row, max_row, min_col, max_col, values_only=False)):
        if i == 0:
            for j, col in enumerate(row):
                head = str(col.value).strip() if col.value != None else None
                if not head:
                    sand_list.append(j + 1)
                elif head == '样品号':
                    head_map['id'] = j + 1
                elif head == '样品名称' or head == '类型' or head == '类别':
                    head_map['sampleName'] = j + 1
                elif head == '黏土基质':
                    head_map['clay'] = j + 1
                elif head == '石英粉砂':
                    head_map['quartz'] = j + 1
                elif head == '砂':
                    sand_list.append(j + 1)
                elif head == '岩屑' or head == '砂或岩屑':
                    head_map['debris'] = j + 1
                elif head == '空洞':
                    head_map['hollow'] = j + 1
                elif head == '其他矿物' and j + 1 == max_col:
                    head_map['other'] = j + 1
                elif head == '石英':
                    head_map['sand']['quartz'] = j + 1
                elif head == '长石':
                    head_map['sand']['feldspar'] = j + 1
                elif head == '其他矿物' or head == '其他':
                    head_map['sand']['Ominerals'] = j + 1
                else:
                    print(f'import_excel ----> 意料之外的head：{head}')
        elif i == 1 and len(sand_list) > 0:
            for j in sand_list:
                head = str(row[j - 1].value).strip() if row[j - 1].value != None else None
                if not head:
                    pass
                elif head == '石英':
                    head_map['sand']['quartz'] = j
                elif head == '长石':
                    head_map['sand']['feldspar'] = j
                elif head == '其他矿物' or head == '其他':
                    head_map['sand']['Ominerals'] = j
                elif head == '小计':
                    pass
                else:
                    print(f'import_excel ----> 意料之外的head：{head}')
        elif isinstance(row[0], MergedCell):
            pass
        else:
            instance = MineContentInfo()
            if head_map['id']:
                v = str(row[head_map['id'] - 1].value).strip()
                instance.id = v if v != 'None' and v != 'n.d.' else None
            ins = MineContentInfo.query.filter_by(id=instance.id).first()
            add_flag = True
            if ins:
                if cover:
                    instance = ins
                    add_flag = False
                    cover_num += 1
                else:
                    continue
            if head_map['sampleName']:
                v = str(row[head_map['sampleName'] - 1].value).strip()
                instance.sampleName = v if v != 'None' and v!= 'n.d.' else None
            if head_map['clay']:
                v = str(row[head_map['clay'] - 1].value).strip()
                instance.clay = float(v) if is_number(v) else None
            if head_map['quartz']:
                v = str(row[head_map['quartz'] - 1].value).strip()
                instance.quartz = float(v) if is_number(v) else None
            if head_map['sand']['quartz']:
                v = str(row[head_map['sand']['quartz'] - 1].value).strip()
                instance.sand['quartz'] = float(v) if is_number(v) else None
            if head_map['sand']['feldspar']:
                v = str(row[head_map['sand']['feldspar'] - 1].value).strip()
                instance.sand['feldspar'] = float(v) if is_number(v) else None
            if head_map['sand']['Ominerals']:
                v = str(row[head_map['sand']['Ominerals'] - 1].value).strip()
                instance.sand['Ominerals'] = float(v) if is_number(v) else None
            if head_map['debris']:
                v = str(row[head_map['debris'] - 1].value).strip()
                instance.debris = float(v) if is_number(v) else None
            if head_map['hollow']:
                v = str(row[head_map['hollow'] - 1].value).strip()
                instance.hollow = float(v) if is_number(v) else None
            if head_map['other']:
                v = str(row[head_map['other'] - 1].value).strip()
                instance.other = float(v) if is_number(v) else None
            if add_flag:
                instance_list[instance.id] = instance
    return instance_list, cover_num

def mine_xrd_process(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int, cover: bool) -> Tuple[Dict[str, MineXRDInfo], int]:
    head_map = {
        "albite": 0,
        "amphibole": 0,
        "analcite": 0,
        "cristobalite": 0,
        "dolomite": 0,
        "hematite": 0,
        "id": 0,
        "magnetite": 0,
        "mica": 0,
        "mullite": 0,
        "potashFeldspar": 0,
        "quartz": 0,
        "tridymite": 0,
        "type": 0
    }
    cover_num = 0
    instance_list = {}
    for i, row in enumerate(ws.iter_rows(min_row, max_row, min_col, max_col, values_only=False)):
        if i == 0:
            for j, col in enumerate(row):
                head = str(col.value).strip() if col.value != None else None
                if not head:
                    pass
                elif head == '样品编号':
                    head_map['id'] = j + 1
                elif head == '类型':
                    head_map['type'] = j + 1
                elif head == '石英':
                    head_map['quartz'] = j + 1
                elif head == '钠长石':
                    head_map['albite'] = j + 1
                elif head == '钾长石':
                    head_map['potashFeldspar'] = j + 1
                elif head == '云母':
                    head_map['mica'] = j + 1
                elif head == '闪石':
                    head_map['amphibole'] = j + 1
                elif head == '赤铁矿':
                    head_map['hematite'] = j + 1
                elif head == '磁铁矿':
                    head_map['magnetite'] = j + 1
                elif head == '白云石':
                    head_map['dolomite'] = j + 1
                elif head == '方沸石':
                    head_map['analcite'] = j + 1
                elif head == '磷石英':
                    head_map['tridymite'] = j + 1
                elif head == '方石英':
                    head_map['cristobalite'] = j + 1
                elif head == '莫来石':
                    head_map['mullite'] = j + 1
                else:
                    print(f'import_excel ----> 意料之外的head：{head}')
        elif isinstance(row[0], MergedCell):
            pass
        else:
            instance = MineXRDInfo()
            if head_map['id']:
                v = str(row[head_map['id'] - 1].value).strip()
                instance.id = v if v != 'None' and v != 'n.d.' else None
            ins = MineXRDInfo.query.filter_by(id=instance.id).first()
            add_flag = True
            if ins:
                if cover:
                    instance = ins
                    add_flag = False
                    cover_num += 1
                else:
                    continue
            if head_map['type']:
                v = str(row[head_map['type'] - 1].value).strip()
                instance.type = v if v != 'None' and v != 'n.d.' else None
            if head_map['quartz']:
                v = str(row[head_map['quartz'] - 1].value).strip()
                instance.quartz = float(v) if is_number(v) else None
            if head_map['albite']:
                v = str(row[head_map['albite'] - 1].value).strip()
                instance.albite = float(v) if is_number(v) else None
            if head_map['potashFeldspar']:
                v = str(row[head_map['potashFeldspar'] - 1].value).strip()
                instance.potashFeldspar = float(v) if is_number(v) else None
            if head_map['mica']:
                v = str(row[head_map['mica'] - 1].value).strip()
                instance.mica = float(v) if is_number(v) else None
            if head_map['amphibole']:
                v = str(row[head_map['amphibole'] - 1].value).strip()
                instance.amphibole = float(v) if is_number(v) else None
            if head_map['hematite']:
                v = str(row[head_map['hematite'] - 1].value).strip()
                instance.hematite = float(v) if is_number(v) else None
            if head_map['magnetite']:
                v = str(row[head_map['magnetite'] - 1].value).strip()
                instance.magnetite = float(v) if is_number(v) else None
            if head_map['dolomite']:
                v = str(row[head_map['dolomite'] - 1].value).strip()
                instance.dolomite = float(v) if is_number(v) else None
            if head_map['analcite']:
                v = str(row[head_map['analcite'] - 1].value).strip()
                instance.analcite = float(v) if is_number(v) else None
            if head_map['tridymite']:
                v = str(row[head_map['tridymite'] - 1].value).strip()
                instance.tridymite = float(v) if is_number(v) else None
            if head_map['cristobalite']:
                v = str(row[head_map['cristobalite'] - 1].value).strip()
                instance.cristobalite = float(v) if is_number(v) else None
            if head_map['mullite']:
                v = str(row[head_map['mullite'] - 1].value).strip()
                instance.mullite = float(v) if is_number(v) else None
            if add_flag:
                instance_list[instance.id] = instance
    return instance_list, cover_num

def mine_chemistry_process(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int, cover: bool) -> Tuple[Dict[str, MineChemistryInfo], int]:
    head_map = {
        'id': 0,
        'type': 0,
        'Na2O': 0,
        'MgO': 0,
        'Al2O3': 0,
        'SiO2': 0,
        'K2O': 0,
        'CaO': 0,
        'Fe2O3': 0
    }
    cover_num = 0
    instance_list = {}
    for i, row in enumerate(ws.iter_rows(min_row, max_row, min_col, max_col, values_only=False)):
        if i == 0:
            for j, col in enumerate(row):
                head = str(col.value).strip() if col.value != None else None
                if not head:
                    pass
                elif head == '样品号':
                    head_map['id'] = j + 1
                elif head == '类型':
                    head_map['type'] = j + 1
                elif head == 'Na2O':
                    head_map['Na2O'] = j + 1
                elif head == 'MgO':
                    head_map['MgO'] = j + 1
                elif head == 'Al2O3':
                    head_map['Al2O3'] = j + 1
                elif head == 'SiO2':
                    head_map['SiO2'] = j + 1
                elif head == 'K2O':
                    head_map['K2O'] = j + 1
                elif head == 'CaO':
                    head_map['CaO'] = j + 1
                elif head == 'Fe2O3':
                    head_map['Fe2O3'] = j + 1
                else:
                    print(f'import_excel ----> 意料之外的head：{head}')
        elif isinstance(row[0], MergedCell):
            pass
        else:
            instance = MineChemistryInfo()
            if head_map['id']:
                v = str(row[head_map['id'] - 1].value).strip()
                instance.id = v if v != 'None' and v != 'n.d.' else None
            ins = MineChemistryInfo.query.filter_by(id=instance.id).first()
            add_flag = True
            if ins:
                if cover:
                    instance = ins
                    add_flag = False
                    cover_num += 1
                else:
                    continue
            if head_map['type']:
                v = str(row[head_map['type'] - 1].value).strip()
                instance.type = v if v != 'None' and v != 'n.d.' else None
            if head_map['Na2O']:
                v = str(row[head_map['Na2O'] - 1].value).strip()
                instance.Na2O = float(v) if is_number(v) else None
            if head_map['MgO']:
                v = str(row[head_map['MgO'] - 1].value).strip()
                instance.MgO = float(v) if is_number(v) else None
            if head_map['Al2O3']:
                v = str(row[head_map['Al2O3'] - 1].value).strip()
                instance.Al2O3 = float(v) if is_number(v) else None
            if head_map['SiO2']:
                v = str(row[head_map['SiO2'] - 1].value).strip()
                instance.SiO2 = float(v) if is_number(v) else None
            if head_map['K2O']:
                v = str(row[head_map['K2O'] - 1].value).strip()
                instance.K2O = float(v) if is_number(v) else None
            if head_map['CaO']:
                v = str(row[head_map['CaO'] - 1].value).strip()
                instance.CaO = float(v) if is_number(v) else None
            if head_map['Fe2O3']:
                v = str(row[head_map['Fe2O3'] - 1].value).strip()
                instance.Fe2O3 = float(v) if is_number(v) else None
            if add_flag:
                instance_list[instance.id] = instance
    return instance_list, cover_num

def mine_thermal_process(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int, cover: bool) -> Tuple[Dict[str, MineThermalInfo], int]:
    head_map = {
        'id': 0,
        'termTemper': 0,
        'fireResis': 0
    }
    cover_num = 0
    instance_list = {}
    for i, row in enumerate(ws.iter_rows(min_row, max_row, min_col, max_col, values_only=False)):
        if i == 0:
            for j, col in enumerate(row):
                head = str(col.value).strip() if col.value != None else None
                if not head:
                    pass
                elif head == '样品号':
                    head_map['id'] = j + 1
                elif head == '终止温度':
                    head_map['termTemper'] = j + 1
                elif head == '耐火度':
                    head_map['fireResis'] = j + 1
                else:
                    print(f'import_excel ----> 意料之外的head：{head}')
        elif isinstance(row[0], MergedCell):
            pass
        else:
            instance = MineThermalInfo()
            if head_map['id']:
                v = str(row[head_map['id'] - 1].value).strip()
                v = re.sub(r'（.*?）', '', v)
                instance.id = v if v != 'None' and v != 'n.d.' else None
            ins = MineThermalInfo.query.filter_by(id=instance.id).first()
            add_flag = True
            if ins:
                if cover:
                    instance = ins
                    add_flag = False
                    cover_num += 1
                else:
                    continue
            if head_map['termTemper']:
                v = str(row[head_map['termTemper'] - 1].value).strip()
                instance.termTemper = float(v) if is_number(v) else None
            if head_map['fireResis']:
                v = str(row[head_map['fireResis'] - 1].value).strip()
                instance.fireResis = float(v) if is_number(v) else None
            if add_flag:
                instance_list[instance.id] = instance
    return instance_list, cover_num

def mine_physics_process(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int, cover: bool) -> Tuple[Dict[str, MinePhysicsInfo], int]:
    head_map = {
        'id': 0,
        'type': 0,
        'apparentPorosity': 0,
        'trueDensity': 0,
        'waterAbsorption': 0
    }
    cover_num = 0
    instance_list = {}
    for i, row in enumerate(ws.iter_rows(min_row, max_row, min_col, max_col, values_only=False)):
        if i == 0:
            for j, col in enumerate(row):
                head = str(col.value).strip() if col.value != None else None
                if not head:
                    pass
                elif head == '样品编号' or head == '样品号':
                    head_map['id'] = j + 1
                elif head == '类型':
                    head_map['type'] = j + 1
                elif head == '显气孔率':
                    head_map['apparentPorosity'] = j + 1
                elif head == '真密度':
                    head_map['trueDensity'] = j + 1
                elif head == '吸水率':
                    head_map['waterAbsorption'] = j + 1
                else:
                    print(f'import_excel ----> 意料之外的head：{head}')
        elif isinstance(row[0], MergedCell):
                pass
        else:
            instance = MinePhysicsInfo()
            if head_map['id']:
                v = str(row[head_map['id'] - 1].value).strip()
                instance.id = v if v != 'None' and v != 'n.d.' else None
            ins = MinePhysicsInfo.query.filter_by(id=instance.id).first()
            add_flag = True
            if ins:
                if cover:
                    instance = ins
                    add_flag = False
                    cover_num += 1
                else:
                    continue
            if head_map['type']:
                v = str(row[head_map['type'] - 1].value).strip()
                instance.type = v if v != 'None' and v != 'n.d.' else None
            if head_map['apparentPorosity']:
                v = str(row[head_map['apparentPorosity'] - 1].value).strip()
                instance.apparentPorosity = float(v) if is_number(v) else None
            if head_map['trueDensity']:
                v = str(row[head_map['trueDensity'] - 1].value).strip()
                instance.trueDensity = float(v) if is_number(v) else None
            if head_map['waterAbsorption']:
                v = str(row[head_map['waterAbsorption'] - 1].value).strip()
                instance.waterAbsorption = float(v) if is_number(v) else None
            if add_flag:
                instance_list[instance.id] = instance
    return instance_list, cover_num

def get_instances(wb: Workbook, cover: bool) -> Tuple[List, List[SampleInfo], int]:
    cover_num = 0
    sample_list = set()
    instance_list = []
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        range_list = get_range(ws)
        if sheetname.strip() == '矿物含量':
            for a, b, c, d in range_list:
                il, con = mine_content_process(ws, a, b, c, d, cover)
                instance_list.extend(il.values())
                sample_list.update(il.keys())
                cover_num += con
        elif sheetname.strip() == '物相（XRD）成分':
            for a, b, c, d in range_list:
                il, con = mine_xrd_process(ws, a, b, c, d, cover)
                instance_list.extend(il.values())
                sample_list.update(il.keys())
                cover_num += con
        elif sheetname.strip() == '化学成分':
            for a, b, c, d in range_list:
                il, con = mine_chemistry_process(ws, a, b, c, d, cover)
                instance_list.extend(il.values())
                sample_list.update(il.keys())
                cover_num += con
        elif sheetname.strip() == '热性能':
            for a, b, c, d in range_list:
                il, con = mine_thermal_process(ws, a, b, c, d, cover)
                instance_list.extend(il.values())
                sample_list.update(il.keys())
                cover_num += con
        elif sheetname.strip() == '物理性能':
            for a, b, c, d in range_list:
                il, con = mine_physics_process(ws, a, b, c, d, cover)
                instance_list.extend(il.values())
                sample_list.update(il.keys())
                cover_num += con
        else:
            print(f'import_excel ----> 意料之外的sheet：{sheetname}')
    sample_list = [SampleInfo(v) for v in sample_list]
    return instance_list, sample_list, cover_num


if __name__ == '__main__':
    print(get_instances(load_workbook('./excels/导入模板.xlsx', read_only=True), cover=True))
